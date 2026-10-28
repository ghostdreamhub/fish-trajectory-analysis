# -*- coding: utf-8 -*-
"""
鱼类性格轨迹指标计算程序 fish_metrics.py

依据《鱼类性格-轨迹指标评价标准_58条扩充版》中“自由游动三维轨迹即可计算”的指标，
从 fish3d 轨迹 CSV 计算每尾鱼的活跃性、空间使用、游动回合和社会距离等指标，并输出数据质量检查。
需要专门实验装置的指标（避难所、新物体、镜子、捕食者模型、双鱼对抗等）不计算，列在输出的“不可计算指标”表中。

用法：
    python fish_metrics.py 轨迹.csv -o 输出文件夹
    python fish_metrics.py 轨迹.csv -o 输出文件夹 -c config.json --tank-json trajectories.json --coverage frames.coverage.json

输入 CSV 必需列：fish_id, segment_id, frame_id, timestamp_ms, x_m, y_m, z_m, point_type
可选列：confidence, covariance_m2, update_kind, smoothed

注意：配置中的速度、距离阈值是占位值（文献值或经验值），换物种、换装置或拿到体长后必须重新标定。
"""
import argparse
import json
import math
import os
import sys
import warnings

import numpy as np
import pandas as pd

# 某些帧只有 1 尾鱼时 nanmin/nanmean 会对全 NaN 切片报警告，结果本身已按 NaN 处理
warnings.filterwarnings("ignore", message="All-NaN slice encountered")
warnings.filterwarnings("ignore", message="Mean of empty slice")

DEFAULT_CONFIG = {
    "tank": {"radius_m": 1.0, "water_depth_m": 0.9},
    "body_length_m": None,
    "point_types_position": ["observed", "interpolated"],
    "point_types_kinematics": ["observed", "interpolated"],
    "max_gap_factor": 3.0,
    "speed_cap_m_s": 1.5,
    "exclude_steps_over_cap": True,
    "v_move_m_s": 0.0175,
    "v_fast_m_s": 0.20,
    "v_freeze_m_s": 0.0175,
    "freeze_min_s": 1.0,
    "wall_zone_width_m": 0.10,
    "center_zone_radius_frac": 0.5,
    "grid_cell_m": 0.10,
    "corrdim": {"r_min_m": 0.01, "r_max_m": 0.30, "n_r": 20, "theiler_s": 1.0,
                "max_points": 2000, "min_points": 200},
    "homebase": {"cell_m": 0.20, "stop_speed_m_s": 0.0175, "stop_min_s": 1.0},
    "social": {"proximity_m": 0.10, "proximity_bl": 1.0, "min_fish_centroid": 3, "nn_max_gap_frames": 2},
    "min_tracked_time_s": 300,
    "min_fps_fine_kinematics": 10,
}

# 输出指标说明：键 -> (中文名, 单位, 对应评价标准序号, 定义/口径)
METRICS = {
    # 数据质量
    "n_points": ("轨迹点数（计算用）", "个", "质量", "按 point_types_position 筛选后的点数"),
    "n_observed": ("实测点数", "个", "质量", "point_type = observed"),
    "n_interpolated": ("插值点数", "个", "质量", "point_type = interpolated（本数据中为单机方位更新 bearing/motion）"),
    "n_predicted_excluded": ("预测点数（已排除）", "个", "质量", "point_type = predicted，默认不参与计算"),
    "n_segments": ("轨迹段数", "段", "质量", "segment_id 个数"),
    "tracked_time_s": ("有效跟踪时长", "s", "质量", "各点代表时长之和（相邻点间隔，超过最大间隔按 1 帧计）"),
    "coverage_pct": ("占记录时长比例", "%", "质量", "有效跟踪时长 / 整段记录时长"),
    "fps_hz": ("帧率", "Hz", "质量", "1 / 帧间隔中位数"),
    "pos_sd_mm_median": ("定位误差中位数", "mm", "质量", "sqrt(协方差矩阵迹)，需 covariance_m2 列"),
    "confidence_median": ("检测置信度中位数", "", "质量", "confidence 列中位数"),
    "n_steps_over_cap": ("超速步数", "步", "质量", "相邻点速度 > speed_cap_m_s，疑似跳点或身份交换"),
    "n_above_water": ("水面以上点数", "个", "质量", "z > 水深"),
    "n_outside_tank": ("缸壁外点数", "个", "质量", "水平半径 > 缸半径"),
    "enough_data": ("跟踪时长是否达标", "", "质量", "有效跟踪时长 ≥ min_tracked_time_s"),
    "warnings": ("警告", "", "质量", "数据量、帧率、体长缺失等提示"),
    # 活跃性
    "speed_median_moving": ("游动速度中位数（运动帧）", "m/s", "#12/#18", "仅取速度 > v_move 的步，Fu 2020 口径"),
    "speed_mean": ("平均速度（全部步）", "m/s", "#18", "总路程 / 有效步时长"),
    "ptm_pct": ("运动时间比 PTM", "%", "#13/#17", "速度 > v_move 的时长 / 有效步时长"),
    "path_length_m": ("总路程", "m", "#14", "有效步（同段、间隔 ≤ 最大间隔、未超速）长度之和"),
    "path_per_min_m": ("每分钟路程", "m/min", "#14", "总路程 / 有效步时长（分钟），便于不同时长比较"),
    "moving_time_s": ("移动时长", "s", "#19", "速度 > v_move 的时长"),
    "turn_rate_median_deg_s": ("转向率中位数", "°/s", "#21", "相邻两步水平运动方向夹角 / 平均步时长，两步速度都 > v_move 时才计"),
    "speed_cv": ("速度变异系数", "", "#45", "SD(速度) / 均值(速度)，全部有效步"),
    "fast_pct": ("快速游动时间比例", "%", "#47", "速度 > v_fast 的时长 / 有效步时长"),
    "accel_mean_abs": ("平均绝对切向加速度", "m/s²", "#44", "相邻两步速度差 / 平均步时长，取绝对值均值；帧率低时不可靠"),
    "freeze_count": ("冻结次数", "次", "#8", "速度持续 < v_freeze 且 ≥ freeze_min_s 的片段数"),
    "freeze_total_s": ("冻结总时长", "s", "#8", "冻结片段时长之和"),
    "freeze_pct": ("冻结时间比例", "%", "#8", "冻结总时长 / 有效步时长"),
    # 游动回合
    "bouts_per_min": ("游动回合数", "次/min", "#20", "连续速度 > v_move 的片段数 / 有效步时长（分钟）"),
    "bout_duration_median_s": ("回合持续时间中位数", "s", "#40", "每个游动回合的时长"),
    "ibi_median_s": ("回合间隔中位数", "s", "#41", "同一连续轨迹中相邻两回合之间的静止时长"),
    "bout_displacement_median_m": ("回合净位移中位数", "m", "#42", "回合起点到终点的直线距离"),
    "bout_peak_speed_median": ("回合峰值速度中位数", "m/s", "#43", "每回合最大速度的中位数"),
    # 空间使用
    "dist_center_mean_m": ("距缸中心平均水平距离", "m", "#46", "sqrt(x²+y²) 的时长加权均值"),
    "dist_wall_mean_m": ("距缸壁平均距离", "m", "#5", "缸半径 − 到中心距离"),
    "wall_zone_pct": ("边缘区时间比例（贴壁）", "%", "#5", "距缸壁 ≤ wall_zone_width_m 的时间比例"),
    "center_zone_pct": ("中心区时间比例", "%", "#5", "到中心距离 ≤ center_zone_radius_frac × 半径 的时间比例"),
    "z_mean_m": ("平均高度（离缸底）", "m", "#7（变体）", "z 的时长加权均值"),
    "bottom_third_pct": ("下 1/3 水层时间比例", "%", "#7（变体）", "z < 水深/3；饲养缸中只能作垂直分布偏好，不等同新异水槽测试"),
    "top_third_pct": ("上 1/3 水层时间比例", "%", "#7（变体）", "z > 2×水深/3"),
    "grid_coverage_pct": ("水平网格到访比例", "%", "#16", "到访过的可进入网格数 / 可进入网格总数（网格边长 grid_cell_m）"),
    "spatial_entropy_norm": ("空间占用熵（归一化）", "", "#39", "H = −Σp·ln p，再除以 ln(可进入网格数)"),
    "corr_dim": ("轨迹相关维数", "", "#38", "Grassberger–Procaccia 法，3D 点，Theiler 窗排除时间近邻"),
    "corr_dim_r2": ("相关维数拟合 R²", "", "#38", "log C(r) 对 log r 线性拟合的决定系数"),
    # 基地行为（饲养缸变体）
    "hb_time_pct": ("偏好区（基地）时间比例", "%", "#34（变体）", "停留时间最长的网格（边长 homebase.cell_m）内的时间比例"),
    "hb_trips": ("完整外出回合数", "次", "#34", "离开基地后再返回基地的完整片段数"),
    "hb_trips_per_min": ("外出回合频率", "次/min", "#34", "完整外出回合数 / 有效跟踪时长（分钟）"),
    "hb_trip_path_mean_m": ("单次外出平均路径长度", "m", "#35", "每个完整外出回合的路径长度均值"),
    "hb_stops_per_trip": ("每次外出停顿数", "次", "#36", "外出回合内速度 < stop_speed 且 ≥ stop_min_s 的片段数均值"),
    # 社会
    "co_present_time_s": ("与其他鱼同时在场时长", "s", "#26", "至少还有 1 尾鱼同帧被定位的时长"),
    "nnd_median_m": ("最近邻距离中位数", "m", "#26", "每帧到最近同伴的 3D 距离"),
    "centroid_dist_mean_m": ("到群体质心平均距离", "m", "#48", "同帧 ≥ min_fish_centroid 尾鱼时，到质心（含自身）的距离均值"),
    "strength": ("加权节点强度", "", "#50", "Σ 与各同伴的接近时间比例 w_ij"),
    "clustering_barrat": ("局部加权聚类系数（Barrat）", "", "#51", "Barrat 等 2004 公式；网络全连通（所有鱼两两接近过）时恒为 1，没有区分度；与 DirectedClustering 包的其他公式不可直接混比"),
    "nn_episode_mean_s": ("最近邻片段平均持续时间", "s", "#52", "同一同伴连续保持为最近邻的时长均值"),
    "strongest_partner_share": ("最强伙伴联结占比", "", "#54", "max_j w_ij / Σ_j w_ij"),
}

PAIR_METRICS = {
    "fish_i": ("鱼 i", "", "", ""),
    "fish_j": ("鱼 j", "", "", ""),
    "co_visible_s": ("同时在场时长", "s", "#49", "两鱼同帧被定位的时长"),
    "dist_median_m": ("两鱼距离中位数", "m", "#49", "同帧 3D 距离"),
    "proximity_m": ("接近判定半径", "m", "#49", "体长已知时为两鱼体长均值 × proximity_bl，否则 proximity_m"),
    "proximity_prop": ("接近时间比例 w_ij", "", "#49", "距离 ≤ 接近半径的时长 / 同时在场时长"),
    "nn_episodes_i_to_j": ("j 作为 i 最近邻的片段数", "次", "#53", "i 的最近邻连续为 j 的片段数"),
    "nn_episode_freq_per_min": ("最近邻片段频率", "次/min", "#53", "片段数 / 同时在场时长（分钟）"),
}

NOT_COMPUTABLE = [
    ("#1–#4、#6", "勇敢性（避难所、模拟捕食者袭击、遮顶）", "实验中没有避难所/活板门、捕食者模型或遮顶，也没有刺激时间戳"),
    ("#10", "明暗箱", "装置不是黑白两半的水箱"),
    ("#11、#15", "首次探索潜伏期、新异物体", "饲养缸不是新环境，没有新物体及其放入时间"),
    ("#22–#25", "社会偏好二元选择、同伴杯、鱼群隔间", "没有隔间式社会性测试装置"),
    ("#27–#29、#58", "镜像测试与好斗分级", "没有镜子；咬击需人工或姿态识别"),
    ("#30、#31", "捕食者检查距离、逃逸启动距离", "没有捕食者模型或标准化逼近的威胁刺激"),
    ("#32", "刺激后逃逸潜伏期", "没有袭击刺激时间戳"),
    ("#55–#57", "双鱼对抗：首次攻击、对抗解决、退让", "没有配对对抗实验；攻击动作需人工标注"),
    ("#7（原定义）", "新异水槽首次进入上层潜伏期", "饲养缸中只能算上/下层时间比例（已作为变体计算）"),
    ("#9", "不规则游动", "原文没有数值阈值，未实现自动判定"),
    ("#26（部分）", "群内前后位置、领导事件", "需要稳定身份与较长连续群体轨迹，未实现"),
    ("#33", "基地形成潜伏期", "原文判据需预先规定，饲养缸无“放入新环境”的起点"),
    ("#37", "基地位置保持率", "需要多次独立测试"),
]


def deep_update(base, new):
    for k, v in new.items():
        if str(k).startswith("_"):
            continue
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            deep_update(base[k], v)
        else:
            base[k] = v
    return base


def load_config(path):
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    if path:
        with open(path, encoding="utf-8") as f:
            deep_update(cfg, json.load(f))
    return cfg


def load_trajectories(path):
    required = ["fish_id", "segment_id", "frame_id", "timestamp_ms", "x_m", "y_m", "z_m", "point_type"]
    df = pd.read_csv(path)
    missing = [c for c in required if c not in df.columns]
    if missing:
        sys.exit(f"轨迹文件缺少必需列：{missing}")
    df["fish_id"] = df["fish_id"].astype(str)
    if "confidence" not in df.columns:
        df["confidence"] = np.nan
    if "covariance_m2" in df.columns:
        def sd(c):
            try:
                return float(np.sqrt(np.trace(np.asarray(json.loads(c), float))))
            except (TypeError, ValueError):
                return np.nan
        df["pos_sd_m"] = df["covariance_m2"].map(sd)
    else:
        df["pos_sd_m"] = np.nan
    return df.sort_values(["fish_id", "segment_id", "frame_id"]).reset_index(drop=True)


def frame_interval_s(df):
    d = df.sort_values(["fish_id", "segment_id", "frame_id"])
    g = d.groupby(["fish_id", "segment_id"])
    dt = g["timestamp_ms"].diff() / g["frame_id"].diff()
    dt = dt[(dt > 0) & np.isfinite(dt)]
    return float(dt.median()) / 1000.0


def body_length(cfg, fid):
    bl = cfg.get("body_length_m")
    if isinstance(bl, dict):
        return bl.get(fid)
    return bl


def point_weights(d, frame_dt, max_dt):
    """每个点代表的时长：到同段下一点的间隔；段末或间隔过大时按 1 帧计。"""
    nxt = d.groupby("segment_id", sort=False)["timestamp_ms"].shift(-1)
    dt = (nxt - d["timestamp_ms"]) / 1000.0
    return dt.where((dt > 0) & (dt <= max_dt), frame_dt).to_numpy()


def add_steps(d, max_dt, cfg):
    """每行记录“本点 → 同段下一点”的一步。"""
    g = d.groupby("segment_id", sort=False)
    dt = (g["timestamp_ms"].shift(-1) - d["timestamp_ms"]) / 1000.0
    dx, dy, dz = (g[c].shift(-1) - d[c] for c in ("x_m", "y_m", "z_m"))
    step = np.sqrt(dx ** 2 + dy ** 2 + dz ** 2)
    speed = step / dt
    valid = dt.notna() & (dt > 0) & (dt <= max_dt)
    over = valid & (speed > cfg["speed_cap_m_s"])
    if cfg["exclude_steps_over_cap"]:
        valid &= ~over
    d = d.copy()
    d["step_dt"] = dt.where(valid)
    d["step_m"] = step.where(valid)
    d["speed"] = speed.where(valid)
    d["speed_h"] = (np.hypot(dx, dy) / dt).where(valid)
    d["heading"] = np.arctan2(dy, dx).where(valid)
    d["valid"] = valid.to_numpy()
    d["over_cap"] = over.to_numpy()
    return d


def true_runs(ok, seg):
    """连续为 True 且在同一轨迹段内的片段，返回 [(起始行, 结束行)]。"""
    runs, s = [], None
    for i in range(len(ok)):
        if ok[i] and s is not None and seg[i] == seg[i - 1]:
            continue
        if s is not None:
            runs.append((s, i - 1))
            s = None
        if ok[i]:
            s = i
    if s is not None:
        runs.append((s, len(ok) - 1))
    return runs


def wsum(w, mask):
    return float(np.sum(w[mask]))


def activity_metrics(k, cfg, fps, warn):
    m = {}
    v, dt, seg = k["speed"].to_numpy(), k["step_dt"].to_numpy(), k["segment_id"].to_numpy()
    valid = k["valid"].to_numpy()
    T = np.nansum(dt)
    if T <= 0:
        warn.append("没有有效的相邻点步，活跃性指标为空")
        return m
    vm, vf, vz = cfg["v_move_m_s"], cfg["v_fast_m_s"], cfg["v_freeze_m_s"]
    moving = valid & (v > vm)
    m["speed_median_moving"] = float(np.nanmedian(v[moving])) if moving.any() else np.nan
    m["path_length_m"] = float(np.nansum(k["step_m"]))
    m["speed_mean"] = m["path_length_m"] / T
    m["ptm_pct"] = 100 * np.nansum(dt[moving]) / T
    m["path_per_min_m"] = m["path_length_m"] / (T / 60)
    m["moving_time_s"] = float(np.nansum(dt[moving]))
    m["fast_pct"] = 100 * np.nansum(dt[valid & (v > vf)]) / T
    vv = v[valid]
    m["speed_cv"] = float(np.std(vv) / np.mean(vv)) if len(vv) > 1 and np.mean(vv) > 0 else np.nan

    # 相邻两步：转向率、加速度
    same = (seg[1:] == seg[:-1]) & valid[1:] & valid[:-1]
    mdt = (dt[1:] + dt[:-1]) / 2
    h = k["heading"].to_numpy()
    vh = k["speed_h"].to_numpy()
    turn_ok = same & (vh[1:] > vm) & (vh[:-1] > vm)
    dtheta = np.abs((h[1:] - h[:-1] + np.pi) % (2 * np.pi) - np.pi)
    rate = np.degrees(dtheta / mdt)
    m["turn_rate_median_deg_s"] = float(np.median(rate[turn_ok])) if turn_ok.any() else np.nan
    acc = np.abs((v[1:] - v[:-1]) / mdt)
    m["accel_mean_abs"] = float(np.mean(acc[same])) if same.any() else np.nan

    # 冻结
    fr = [(s, e) for s, e in true_runs(valid & (v < vz), seg) if np.nansum(dt[s:e + 1]) >= cfg["freeze_min_s"]]
    m["freeze_count"] = len(fr)
    m["freeze_total_s"] = float(sum(np.nansum(dt[s:e + 1]) for s, e in fr))
    m["freeze_pct"] = 100 * m["freeze_total_s"] / T

    # 游动回合
    bouts = true_runs(moving, seg)
    m["bouts_per_min"] = len(bouts) / (T / 60)
    if bouts:
        x, y, z = (k[c].to_numpy() for c in ("x_m", "y_m", "z_m"))
        dur = [np.nansum(dt[s:e + 1]) for s, e in bouts]
        # 回合终点 = 最后一步的下一点
        disp = [math.dist((x[s], y[s], z[s]), (x[e + 1], y[e + 1], z[e + 1])) for s, e in bouts if e + 1 < len(x)]
        peak = [np.nanmax(v[s:e + 1]) for s, e in bouts]
        m["bout_duration_median_s"] = float(np.median(dur))
        m["bout_displacement_median_m"] = float(np.median(disp)) if disp else np.nan
        m["bout_peak_speed_median"] = float(np.median(peak))
        ibi = []
        for (s0, e0), (s1, e1) in zip(bouts[:-1], bouts[1:]):
            gap = slice(e0 + 1, s1)
            if seg[s1] == seg[e0] and valid[gap].all():
                ibi.append(np.nansum(dt[gap]))
        m["ibi_median_s"] = float(np.median(ibi)) if ibi else np.nan
    if fps < cfg["min_fps_fine_kinematics"]:
        warn.append(f"帧率 {fps:.1f} Hz < {cfg['min_fps_fine_kinematics']} Hz，回合、加速度、转向率不可靠")
    return m


def grid_occupancy(x, y, w, R, cell):
    n = max(1, int(round(2 * R / cell)))
    edges = np.linspace(-R, R, n + 1)
    c = (edges[:-1] + edges[1:]) / 2
    cx, cy = np.meshgrid(c, c, indexing="ij")
    acc = np.hypot(cx, cy) <= R
    ix = np.clip(np.digitize(x, edges) - 1, 0, n - 1)
    iy = np.clip(np.digitize(y, edges) - 1, 0, n - 1)
    occ = np.zeros((n, n))
    np.add.at(occ, (ix, iy), w)
    return occ, acc, ix, iy


def correlation_dimension(P, t, cc, warn):
    if len(P) < cc["min_points"]:
        warn.append(f"相关维数：点数 {len(P)} < {cc['min_points']}，未计算")
        return np.nan, np.nan
    if len(P) > cc["max_points"]:
        idx = np.linspace(0, len(P) - 1, cc["max_points"]).astype(int)
        P, t = P[idx], t[idx]
    # NumPy equivalent of condensed pairwise distances; avoids a SciPy runtime dependency.
    ii, jj = np.triu_indices(len(P), 1)
    temporal = np.abs(t[ii] - t[jj]) >= cc["theiler_s"]
    D = np.linalg.norm(P[ii[temporal]] - P[jj[temporal]], axis=1)
    if len(D) < 100:
        warn.append("相关维数：排除时间近邻后点对太少，未计算")
        return np.nan, np.nan
    D.sort()
    r = np.logspace(math.log10(cc["r_min_m"]), math.log10(cc["r_max_m"]), cc["n_r"])
    C = np.searchsorted(D, r) / len(D)
    ok = (C > 0) & (C < 0.9)
    if ok.sum() < 4:
        warn.append("相关维数：有效尺度点不足 4 个，未计算")
        return np.nan, np.nan
    lx, ly = np.log(r[ok]), np.log(C[ok])
    slope, icpt = np.polyfit(lx, ly, 1)
    pred = slope * lx + icpt
    r2 = 1 - np.sum((ly - pred) ** 2) / np.sum((ly - ly.mean()) ** 2)
    return float(slope), float(r2)


def spatial_metrics(p, w, cfg, warn):
    R, depth = cfg["tank"]["radius_m"], cfg["tank"]["water_depth_m"]
    x, y, z = (p[c].to_numpy() for c in ("x_m", "y_m", "z_m"))
    r = np.hypot(x, y)
    W = w.sum()
    m = {
        "dist_center_mean_m": float(np.sum(r * w) / W),
        "dist_wall_mean_m": float(np.sum((R - r) * w) / W),
        "wall_zone_pct": 100 * wsum(w, r >= R - cfg["wall_zone_width_m"]) / W,
        "center_zone_pct": 100 * wsum(w, r <= cfg["center_zone_radius_frac"] * R) / W,
        "z_mean_m": float(np.sum(z * w) / W),
        "bottom_third_pct": 100 * wsum(w, z < depth / 3) / W,
        "top_third_pct": 100 * wsum(w, z > 2 * depth / 3) / W,
    }
    occ, acc, _, _ = grid_occupancy(x, y, w, R, cfg["grid_cell_m"])
    oa = occ[acc]
    K = int(acc.sum())
    m["grid_coverage_pct"] = 100 * (oa > 0).sum() / K
    pk = oa[oa > 0] / oa.sum()
    m["spatial_entropy_norm"] = float(-(pk * np.log(pk)).sum() / np.log(K)) if K > 1 else np.nan
    t = (p["timestamp_ms"].to_numpy() - p["timestamp_ms"].min()) / 1000.0
    m["corr_dim"], m["corr_dim_r2"] = correlation_dimension(np.c_[x, y, z], t, cfg["corrdim"], warn)
    return m


def homebase_metrics(k, w, cfg, T_track):
    hb = cfg["homebase"]
    R = cfg["tank"]["radius_m"]
    x, y = k["x_m"].to_numpy(), k["y_m"].to_numpy()
    occ, acc, ix, iy = grid_occupancy(x, y, w, R, hb["cell_m"])
    occ[~acc] = 0
    bi, bj = np.unravel_index(np.argmax(occ), occ.shape)
    inside = (ix == bi) & (iy == bj)
    m = {"hb_time_pct": 100 * wsum(w, inside) / w.sum()}
    seg, valid = k["segment_id"].to_numpy(), k["valid"].to_numpy()
    v, dt, step = k["speed"].to_numpy(), k["step_dt"].to_numpy(), k["step_m"].to_numpy()
    paths, stops = [], []
    for s, e in true_runs(~inside, seg):
        # 完整外出：前一点在基地、后一点在基地、同段且中间步全部有效
        if s == 0 or e + 1 >= len(inside) or not (inside[s - 1] and inside[e + 1]):
            continue
        if seg[s - 1] != seg[s] or seg[e + 1] != seg[e]:
            continue
        steps = slice(s - 1, e + 1)
        if not valid[steps].all():
            continue
        paths.append(np.nansum(step[steps]))
        sub_ok = valid[steps] & (v[steps] < hb["stop_speed_m_s"])
        sub_dt = dt[steps]
        stops.append(sum(1 for a, b in true_runs(sub_ok, seg[steps]) if np.nansum(sub_dt[a:b + 1]) >= hb["stop_min_s"]))
    m["hb_trips"] = len(paths)
    m["hb_trips_per_min"] = len(paths) / (T_track / 60) if T_track > 0 else np.nan
    m["hb_trip_path_mean_m"] = float(np.mean(paths)) if paths else np.nan
    m["hb_stops_per_trip"] = float(np.mean(stops)) if stops else np.nan
    return m


def barrat_clustering(Wm):
    n = len(Wm)
    A = (Wm > 0).astype(float)
    out = np.full(n, np.nan)
    for i in range(n):
        k, s = A[i].sum(), Wm[i].sum()
        if k < 2 or s <= 0:
            continue
        tot = 0.0
        for j in range(n):
            for h in range(n):
                if j != h and j != i and h != i:
                    tot += (Wm[i, j] + Wm[i, h]) / 2 * A[i, j] * A[i, h] * A[j, h]
        out[i] = tot / (s * (k - 1))
    return out


def social_metrics(pos, kin, ids, frame_dt, cfg):
    sc = cfg["social"]
    n = len(ids)
    piv = pos.pivot_table(index="frame_id", columns="fish_id", values=["x_m", "y_m", "z_m"], aggfunc="first")
    frames = piv.index.to_numpy()
    X = np.stack([np.stack([piv[c][f].to_numpy() if f in piv[c] else np.full(len(frames), np.nan)
                            for f in ids], axis=1) for c in ("x_m", "y_m", "z_m")], axis=2)
    present = ~np.isnan(X[:, :, 0])
    D = np.linalg.norm(X[:, :, None, :] - X[:, None, :, :], axis=3)
    D[:, np.arange(n), np.arange(n)] = np.nan
    cnt = present.sum(1)

    # 接近半径
    thr = np.full((n, n), sc["proximity_m"], float)
    bls = [body_length(cfg, f) for f in ids]
    for i in range(n):
        for j in range(n):
            if bls[i] and bls[j]:
                thr[i, j] = (bls[i] + bls[j]) / 2 * sc["proximity_bl"]
    both = present[:, :, None] & present[:, None, :]
    co = both.sum(0).astype(float)
    np.fill_diagonal(co, 0)
    prox = (both & (D <= thr)).sum(0).astype(float)
    Wm = np.divide(prox, co, out=np.zeros_like(prox), where=co > 0)
    np.fill_diagonal(Wm, 0)

    with np.errstate(all="ignore"):
        nnd = np.nanmin(D, axis=2)
    has_nb = present & (cnt[:, None] >= 2)
    nn_idx = np.where(has_nb, np.nanargmin(np.where(np.isnan(D), np.inf, D), axis=2), -1)

    big = cnt >= sc["min_fish_centroid"]
    cen = np.nanmean(np.where(present[:, :, None], X, np.nan), axis=1)
    cdist = np.linalg.norm(X - cen[:, None, :], axis=2)

    per, pairs = {}, []
    episodes = np.zeros((n, n))
    clus = barrat_clustering(Wm)
    for i, f in enumerate(ids):
        sel = has_nb[:, i]
        eps, cur, start, last = [], None, None, None
        for fr, j in zip(frames[sel], nn_idx[sel, i]):
            if cur is not None and j == cur and fr - last <= sc["nn_max_gap_frames"] + 1:
                last = fr
                continue
            if cur is not None:
                eps.append((cur, (last - start + 1) * frame_dt))
                episodes[i, cur] += 1
            cur, start, last = j, fr, fr
        if cur is not None:
            eps.append((cur, (last - start + 1) * frame_dt))
            episodes[i, cur] += 1
        s = Wm[i].sum()
        per[f] = {
            "co_present_time_s": float(sel.sum() * frame_dt),
            "nnd_median_m": float(np.nanmedian(nnd[sel, i])) if sel.any() else np.nan,
            "centroid_dist_mean_m": float(np.nanmean(cdist[big & present[:, i], i])) if (big & present[:, i]).any() else np.nan,
            "strength": float(s),
            "clustering_barrat": float(clus[i]),
            "nn_episode_mean_s": float(np.mean([d for _, d in eps])) if eps else np.nan,
            "strongest_partner_share": float(Wm[i].max() / s) if s > 0 else np.nan,
        }
    for i in range(n):
        for j in range(n):
            if i == j or co[i, j] == 0:
                continue
            dd = D[both[:, i, j], i, j]
            pairs.append({
                "fish_i": ids[i], "fish_j": ids[j],
                "co_visible_s": co[i, j] * frame_dt,
                "dist_median_m": float(np.median(dd)),
                "proximity_m": thr[i, j],
                "proximity_prop": Wm[i, j],
                "nn_episodes_i_to_j": int(episodes[i, j]),
                "nn_episode_freq_per_min": episodes[i, j] / (co[i, j] * frame_dt / 60),
            })

    # 群体逐帧：个体间平均距离、方向一致性（极化度）
    iu = np.triu_indices(n, 1)
    with np.errstate(all="ignore"):
        iid = np.nanmean(D[:, iu[0], iu[1]], axis=1)
    mv = kin[kin["valid"] & (kin["speed_h"] > cfg["v_move_m_s"])]
    hp = mv.pivot_table(index="frame_id", columns="fish_id", values="heading", aggfunc="first").reindex(frames)
    U = np.exp(1j * hp.to_numpy())
    nm = np.sum(~np.isnan(U), axis=1)
    with np.errstate(all="ignore"):
        pol = np.where(nm >= 2, np.abs(np.nansum(U, axis=1)) / nm, np.nan)
    group = pd.DataFrame({"frame_id": frames, "n_present": cnt, "iid_mean_m": np.where(cnt >= 2, iid, np.nan),
                          "n_moving": nm, "polarization": pol})
    return per, pd.DataFrame(pairs), group


def fish_metrics(df_all, fid, frame_dt, t_span, cfg):
    warn = []
    d = df_all[df_all["fish_id"] == fid]
    max_dt = cfg["max_gap_factor"] * frame_dt
    p = d[d["point_type"].isin(cfg["point_types_position"])].reset_index(drop=True)
    k = d[d["point_type"].isin(cfg["point_types_kinematics"])].reset_index(drop=True)
    m = {"n_points": len(p),
         "n_observed": int((d["point_type"] == "observed").sum()),
         "n_interpolated": int((d["point_type"] == "interpolated").sum()),
         "n_predicted_excluded": int(((d["point_type"] == "predicted") & ~d["point_type"].isin(cfg["point_types_position"])).sum()),
         "n_segments": int(d["segment_id"].nunique())}
    if len(p) < 3:
        m["warnings"] = "点数不足 3，未计算"
        return m, k
    w = point_weights(p, frame_dt, max_dt)
    T_track = float(w.sum())
    R, depth = cfg["tank"]["radius_m"], cfg["tank"]["water_depth_m"]
    m.update({
        "tracked_time_s": T_track,
        "coverage_pct": 100 * T_track / t_span if t_span > 0 else np.nan,
        "fps_hz": 1 / frame_dt,
        "pos_sd_mm_median": float(p["pos_sd_m"].median() * 1000) if p["pos_sd_m"].notna().any() else np.nan,
        "confidence_median": float(p["confidence"].median()),
        "n_above_water": int((p["z_m"] > depth).sum()),
        "n_outside_tank": int((np.hypot(p["x_m"], p["y_m"]) > R).sum()),
        "enough_data": "是" if T_track >= cfg["min_tracked_time_s"] else "否",
    })
    if T_track < cfg["min_tracked_time_s"]:
        warn.append(f"有效跟踪 {T_track:.0f} s < {cfg['min_tracked_time_s']} s，结果只能用于测试程序")
    if body_length(cfg, fid) is None:
        warn.append("未提供体长，距离阈值使用绝对值")
    k = add_steps(k, max_dt, cfg)
    m["n_steps_over_cap"] = int(k["over_cap"].sum())
    m.update(activity_metrics(k, cfg, 1 / frame_dt, warn))
    m.update(spatial_metrics(p, w, cfg, warn))
    wk = point_weights(k, frame_dt, max_dt)
    m.update(homebase_metrics(k, wk, cfg, float(wk.sum())))
    m["warnings"] = "；".join(warn)
    return m, k


def compute_all(df, cfg):
    """计算全部个体、成对和群体逐帧指标；供命令行和其他分析脚本调用。"""
    ids = sorted(df["fish_id"].unique())
    frame_dt = frame_interval_s(df)
    t_span = (df["timestamp_ms"].max() - df["timestamp_ms"].min()) / 1000.0 + frame_dt
    rows, kin = [], []
    for f in ids:
        m, k = fish_metrics(df, f, frame_dt, t_span, cfg)
        rows.append({"fish_id": f, **m})
        kin.append(k)
    per = pd.DataFrame(rows)
    pos_all = df[df["point_type"].isin(cfg["point_types_position"])]
    kin_all = pd.concat([k for k in kin if "valid" in k.columns], ignore_index=True)
    soc, pairs, group = social_metrics(pos_all, kin_all, ids, frame_dt, cfg)
    for c in next(iter(soc.values())).keys():
        per[c] = per["fish_id"].map(lambda f: soc[f][c])
    order = ["fish_id"] + [c for c in METRICS if c in per.columns and c != "warnings"] + ["warnings"]
    return per[order], pairs, group, ids, frame_dt, t_span


def labelled(df, spec):
    cols = {}
    for c in df.columns:
        if c in spec:
            name, unit, no, _ = spec[c]
            cols[c] = f"{name}（{unit}）{no and '【' + no + '】'}" if unit else f"{name}{no and '【' + no + '】'}"
    return df.rename(columns=cols)


def write_excel(path, sheets):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E5A")
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0000"
            for i, col in enumerate(frame.columns, 1):
                vals = [str(col)] + [str(v) for v in frame[col].head(50)]
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(v) for v in vals) * 1.2))
            ws.row_dimensions[1].height = 45
            ws.freeze_panes = "B2"


def plot_qc(df, ids, cfg, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial"]
    plt.rcParams["axes.unicode_minus"] = False
    R, depth = cfg["tank"]["radius_m"], cfg["tank"]["water_depth_m"]
    fig, ax = plt.subplots(1, 3, figsize=(16, 5))
    t0 = df["timestamp_ms"].min()
    for f in ids:
        d = df[df["fish_id"] == f]
        for _, s in d.groupby("segment_id"):
            ax[0].plot(s["x_m"], s["y_m"], lw=0.8, color=f"C{ids.index(f)}")
            ax[1].plot(np.hypot(s["x_m"], s["y_m"]), s["z_m"], lw=0.8, color=f"C{ids.index(f)}")
            ax[2].plot((s["timestamp_ms"] - t0) / 1000, s["z_m"], lw=0.8, color=f"C{ids.index(f)}")
        ax[0].plot([], [], color=f"C{ids.index(f)}", label=f)
    th = np.linspace(0, 2 * np.pi, 200)
    ax[0].plot(R * np.cos(th), R * np.sin(th), "k--", lw=0.8)
    ax[0].set(aspect="equal", xlabel="x (m)", ylabel="y (m)", title="俯视轨迹（虚线为缸壁）")
    ax[0].legend(fontsize=8)
    ax[1].axhline(depth, color="b", ls="--", lw=0.8)
    ax[1].axvline(R, color="k", ls="--", lw=0.8)
    ax[1].set(xlabel="到中心距离 (m)", ylabel="z (m)", title="侧视（蓝线水面，黑线缸壁）")
    ax[2].axhline(depth, color="b", ls="--", lw=0.8)
    ax[2].set(xlabel="时间 (s)", ylabel="z (m)", title="高度随时间")
    fig.tight_layout()
    fig.savefig(path, dpi=130)
    plt.close(fig)


def main():
    ap = argparse.ArgumentParser(description="鱼类性格轨迹指标计算")
    ap.add_argument("csv", help="轨迹 CSV（fish3d 格式）")
    ap.add_argument("-o", "--out", required=True, help="输出文件夹")
    ap.add_argument("-c", "--config", help="参数配置 JSON")
    ap.add_argument("--tank-json", help="含 tank 字段的 trajectories.json 或 运行参数.json，用于读取缸半径和水深")
    ap.add_argument("--coverage", help="frames.coverage.json，可选，汇总到质量检查")
    a = ap.parse_args()

    cfg = load_config(a.config)
    if a.tank_json:
        with open(a.tank_json, encoding="utf-8") as f:
            tank = json.load(f).get("tank", {})
        for key in ("radius_m", "water_depth_m"):
            if key in tank:
                cfg["tank"][key] = tank[key]
    os.makedirs(a.out, exist_ok=True)

    df = load_trajectories(a.csv)
    per, pairs, group, ids, frame_dt, t_span = compute_all(df, cfg)
    pos_all = df[df["point_type"].isin(cfg["point_types_position"])]

    qc_global = [
        ("轨迹文件", os.path.abspath(a.csv)),
        ("记录起止（北京时间）", f"{pd.to_datetime(df.timestamp_ms.min(), unit='ms', utc=True).tz_convert('Asia/Shanghai')} → "
                             f"{pd.to_datetime(df.timestamp_ms.max(), unit='ms', utc=True).tz_convert('Asia/Shanghai')}"),
        ("记录时长 (s)", round(t_span, 2)),
        ("帧间隔 (s) / 帧率 (Hz)", f"{frame_dt:.4f} / {1 / frame_dt:.2f}"),
        ("鱼 ID 数", len(ids)),
        ("同帧在场鱼数分布（帧数）", json.dumps(group["n_present"].value_counts().sort_index().to_dict())),
        ("个体间平均距离中位数 (m)【#26】", round(float(group["iid_mean_m"].median()), 4)),
        ("极化度中位数【#26】", round(float(group["polarization"].median()), 4) if group["polarization"].notna().any() else "无（同帧运动鱼不足 2 尾）"),
        ("缸半径 / 水深 (m)", f"{cfg['tank']['radius_m']} / {cfg['tank']['water_depth_m']}"),
    ]
    if a.coverage:
        with open(a.coverage, encoding="utf-8") as f:
            cov = json.load(f)
        for key, label in [("expected_fish", "配置鱼数 expected_fish"),
                           ("true_3d_observed_fish_frame_rate", "真实三维观测覆盖率"),
                           ("usable_fish_frame_rate", "可用覆盖率（含插值/预测）"),
                           ("two_camera_effective_frame_rate", "双机有效帧比例")]:
            if key in cov:
                qc_global.append((label, cov[key]))
        for cam, c in cov.get("camera_detection_counts", {}).items():
            qc_global.append((f"{cam} 检出数 / 每帧", f"{c.get('detected')} / {c.get('detected', 0) / max(cov.get('frames', 1), 1):.2f}"))

    spec_rows = [{"输出列": f"{v[0]}", "单位": v[1], "对应评价标准": v[2], "定义 / 口径": v[3]} for v in {**METRICS, **PAIR_METRICS}.values() if v[2]]
    params = pd.DataFrame([{"参数": k, "取值": json.dumps(v, ensure_ascii=False)} for k, v in cfg.items()])
    sheets = {
        "个体指标": labelled(per, METRICS),
        "成对指标": labelled(pairs, PAIR_METRICS) if len(pairs) else pd.DataFrame({"说明": ["同帧在场的鱼不足 2 尾"]}),
        "群体逐帧": group,
        "整体质量": pd.DataFrame(qc_global, columns=["项目", "值"]),
        "参数": params,
        "指标说明": pd.DataFrame(spec_rows),
        "不可计算指标": pd.DataFrame(NOT_COMPUTABLE, columns=["评价标准序号", "指标", "原因"]),
    }
    xlsx = os.path.join(a.out, "指标结果.xlsx")
    write_excel(xlsx, sheets)
    per.to_csv(os.path.join(a.out, "individual_metrics.csv"), index=False, encoding="utf-8-sig")
    if len(pairs):
        pairs.to_csv(os.path.join(a.out, "pair_metrics.csv"), index=False, encoding="utf-8-sig")
    group.to_csv(os.path.join(a.out, "group_frames.csv"), index=False, encoding="utf-8-sig")
    try:
        plot_qc(pos_all, ids, cfg, os.path.join(a.out, "轨迹检查图.png"))
    except ImportError:
        print("未安装 matplotlib，跳过检查图")
    print(f"完成：{len(ids)} 尾鱼，记录 {t_span:.1f} s，帧率 {1 / frame_dt:.2f} Hz")
    print(f"结果：{xlsx}")
    for _, r in per.iterrows():
        if r.get("warnings"):
            print(f"  {r['fish_id']}：{r['warnings']}")


if __name__ == "__main__":
    main()
